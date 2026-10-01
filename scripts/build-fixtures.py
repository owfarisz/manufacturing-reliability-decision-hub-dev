"""Read-only, deterministic workbook adapter. Never saves an input workbook."""
from pathlib import Path
import os, json, hashlib, re
from collections import Counter
import openpyxl
BASE=Path(__file__).resolve().parents[1]
ROOT=Path(os.environ.get('CASE2_DATA_ROOT',str(BASE.parent/'Case 2_ Intelligence Manufacturing'))).resolve()
OUT=BASE/'src/generated'; OUT.mkdir(parents=True,exist_ok=True)
def clean(x):
 return x.isoformat(sep=' ') if hasattr(x,'hour') else x.isoformat() if hasattr(x,'isoformat') else x
def rows(w,s): return [[clean(v) for v in r] for r in w[s].values]
def records(r,h=0): return [dict(zip(r[h],x)) for x in r[h+1:] if any(v is not None for v in x)]
def prov(t,p,s,f,**kw): return dict(type=t,sourceFile=str(p.relative_to(ROOT)),sheet=s,sourceField=f,**kw)
data={'assets':{},'registry':{},'validation':[],'incidents':[],'sourceHashes':{}}
for p in sorted(ROOT.rglob('*.xlsx')):
 before=hashlib.sha256(p.read_bytes()).hexdigest(); w=openpyxl.load_workbook(p,read_only=True,data_only=True)
 if p.parent.name=='Production Data':
  tags=records(rows(w,'PI Tag')) # MUST precede measurements.
  rr=rows(w,'Sheet2'); rec=records(rr); assert len(rec)==720
  tag=re.search(r'RCA\d (.+)\.xlsx',p.name)[1]; a=data['assets'].setdefault(tag,{})
  a.update(hourly=rec,tags=tags,hourlySource=prov('ACTUAL_HOURLY',p,'Sheet2','All original fields',timestampOrRange=f"{rec[0]['Timestamp']} to {rec[-1]['Timestamp']}",limitation='Local source timestamps; timezone unspecified. Historical snapshot, not live.'))
  for field in rr[0]: data['registry'][tag+':hourly:'+field]=prov('ACTUAL_HOURLY',p,'Sheet2',field,unit=next((t['engunits'] for t in tags if t['Name']==field),'local timestamp'))
  off=[r['Timestamp'] for r in rec if r['RUN_STATUS']=='OFF']; a['off']={'count':len(off),'start':off[0] if off else None,'end':off[-1] if off else None}
  data['validation'].append(dict(file=p.name,sheet='Sheet2',rows=len(rec),start=rec[0]['Timestamp'],end=rec[-1]['Timestamp'],nulls={f:sum(r[f] is None for r in rec) for f in rr[0]},readOrder=['PI Tag','Sheet2']))
 elif p.parent.name=='Equipment Performance':
  info=rows(w,'Equipment Info'); hist=rows(w,'Condition History'); summary=rows(w,'Performance Summary'); rec=records(hist); assert len(rec)==26
  tag=info[3][1]; a=data['assets'].setdefault(tag,{})
  a.update(info=info,weekly=rec,summary=summary,fields=hist[0][2:6],weeklySource=prov('ACTUAL_WEEKLY',p,'Condition History','All condition fields',timestampOrRange=f"{rec[0]['Date']} to {rec[-1]['Date']}",limitation='26 weekly readings and one failure. No sub-week temporal precision or guaranteed early-warning lead time.'))
  for field in hist[0]: data['registry'][tag+':weekly:'+field]=prov('ACTUAL_WEEKLY',p,'Condition History',field)
  a['thresholds']=[{'field':r[2],'alarm':float(r[3].split('/')[0]),'trip':float(r[3].split('/')[1]),'provenance':prov('ACTUAL_WEEKLY',p,'Equipment Info',r[2],limitation='Source limits, demo configuration v1; not approved for plant operation.')} for r in info[4:8]]
  data['validation'].append(dict(file=p.name,sheet='Condition History',rows=26,start=rec[0]['Date'],end=rec[-1]['Date'],nulls={f:sum(r[f] is None for r in rec) for f in hist[0]},sheetsRead=w.sheetnames))
 else:
  data['dashboard']=rows(w,'Dashboard'); rr=rows(w,'Incident Database'); rec=records(rr,2); assert len(rec)==380
  data['incidents']=rec; data['incidentSource']=prov('INCIDENT_RECORDED',p,'Incident Database','Rows 4–383',limitation='Recorded portfolio snapshot. Status is not a live plant status.')
  for field in rr[2]: data['registry']['incidents:'+field]=prov('INCIDENT_RECORDED',p,'Incident Database',field)
  data['validation'].append(dict(file=p.name,sheet='Incident Database',rows=380,headerRow=3,nulls={f:sum(r[f] is None for r in rec) for f in rr[2]},start=min(r['Date of Occur.'] for r in rec),end=max(r['Date of Occur.'] for r in rec),sheetsRead=w.sheetnames))
 w.close(); assert before==hashlib.sha256(p.read_bytes()).hexdigest(); data['sourceHashes'][str(p.relative_to(ROOT))]=before
r=data['incidents']; data['portfolio']={'count':len(r),'downtime':round(sum(x['Downtime (hrs)'] for x in r),2),'actualLoss':round(sum(x['Act. Loss (k US$)'] for x in r),2),'potentialLoss':round(sum(x['Pot. Loss (k US$)'] for x in r),2),'totalLoss':round(sum(x['Total Loss (k US$)'] for x in r),2),'statuses':dict(Counter(x['Overall Status'] for x in r)),'families':[]}
for name,pattern in [('Leakage-related','leak'),('Vibration-related','vibration'),('Fouling','foul'),('Bearing','bearing')]:
 group=[x for x in r if pattern in x['Risk Case Title'].lower()]; data['portfolio']['families'].append({'name':name,'count':len(group),'actualLoss':round(sum(x['Act. Loss (k US$)'] for x in group),2),'pattern':pattern})
data['conflicts']=[
 {'title':'KO vibration unit definition unresolved — engineering validation required','detail':'KO3201_VIB is MM/S in PI Tag; weekly DE radial vibration is micron. No conversion, joining, or micron threshold applied to hourly data.','source':'Production Data - RCA2 KO-3201.xlsx / PI Tag vs Equipment Performance / Condition History'},
 {'title':'KO oil-pressure evidence contradicts across sources','detail':'RCA2 slide 6 reports normal 1.8 barg. Weekly failure row reports 1.078 barg. Retain both; validate instrument and sampling context.','source':'RCA2 slide 6 / Equipment Performance KO Condition History'},
 {'title':'KO vibration thresholds have different contexts','detail':'RCA2 slide 7 describes historical 60 micron alert; slides 5–6 and Equipment Info specify 45 micron. Prototype uses 45 micron only on weekly displacement, as unapproved demo configuration.','source':'RCA2 slides 5–7 / Equipment Info'},
 {'title':'HE discharge pressure is not tube-side differential pressure','detail':'HE3301_DISP remains generic discharge pressure (BARG). Fouling uses weekly Tube-side dP (bar).','source':'Production Data - RCA4 HE-3301.xlsx / PI Tag'},
 {'title':'HE 13 OFF samples versus 12 hours elapsed downtime','detail':'Hourly OFF rows span 21-May 09:00 through 21:00 inclusive (13 samples). RCA reports 09:00–21:00 = 12 elapsed hours. Do not equate samples with duration.','source':'HE Sheet2 / RCA4 slide 3'},
 {'title':'Weekly records cannot be fully rate-normalized from supplied hourly flow','detail':'Weekly timestamps are date-only and extend outside the hourly month. Normalized demonstrations use explicit assumed rate, never a fabricated actual join.','source':'HE Condition History / Sheet2'},
 {'title':'KO weekly contamination precedes RCA narrative window','detail':'Weekly water first exceeds 500 ppm on 11-Feb; RCA describes approximately two weeks pre-trip. Historical temporal precision differs; no guaranteed warning claim.','source':'KO Condition History / RCA2 slide 4'},
 {'title':'Equipment nameplate template fields require validation','detail':'HE Equipment Info includes bearing/seal design-life text and generic monitoring method. Do not use these template fields as verified exchanger design.','source':'Equipment Performance HE / Equipment Info'}]
ui={'assets':{k:{**{f:v for f,v in a.items() if f!='hourly'},'hourly':a['hourly'][::6]+[r for r in a['hourly'] if r['RUN_STATUS']=='OFF']} for k,a in data['assets'].items() if k in ('KO-3201','HE-3301')},'portfolio':data['portfolio'],'conflicts':data['conflicts'],'registry':{k:v for k,v in data['registry'].items() if k.startswith(('KO-3201','HE-3301','incidents:'))},'incidentSource':data['incidentSource']}
(OUT/'ui.json').write_text(json.dumps(ui,ensure_ascii=False,separators=(',',':')))
(OUT/'fixtures.json').write_text(json.dumps(data,ensure_ascii=False,separators=(',',':')))
(BASE/'docs/data-validation.json').write_text(json.dumps({'files':data['validation'],'conflicts':data['conflicts'],'portfolio':data['portfolio'],'sourceHashes':data['sourceHashes']},indent=2))
print(json.dumps({'workbooks':len(data['validation']),'hourly':sum(len(a['hourly']) for a in data['assets'].values()),'weekly':sum(len(a['weekly']) for a in data['assets'].values()),'portfolio':data['portfolio'],'sourceUnchanged':True},indent=2))
